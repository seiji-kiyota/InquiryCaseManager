import csv
import io
from datetime import date, datetime, time

from openpyxl import load_workbook

from common.case_service import create_case_from_inquiry, search_cases
from common.db import init_db
from common.export_service import (
    CASE_SHEET_NAME,
    INQUIRY_SHEET_NAME,
    OVERDUE_LABEL,
    case_export_headers,
    create_case_csv_bytes,
    create_case_excel_bytes,
    create_inquiry_csv_bytes,
    create_inquiry_excel_bytes,
    generate_export_filename,
    inquiry_export_headers,
    prepare_case_export_rows,
    prepare_inquiry_export_rows,
)
from common.inquiry_service import create_inquiry, search_inquiries


TODAY = date(2026, 9, 21)


def _inquiry(**overrides):
    payload = {
        "received_date": date(2026, 9, 10),
        "received_time": time(10, 0, 0),
        "customer_name": "山田太郎",
        "company_name": "株式会社サンプル",
        "phone": "03-1111-1111",
        "email": "taro@example.com",
        "channel": "電話",
        "category": "商品",
        "subject": "見積依頼",
        "description": "商品の見積をお願いします。",
        "priority": "高",
        "assignee": "山田",
        "due_date": date(2026, 9, 30),
        "status": "未対応",
        "notes": "",
    }
    payload.update(overrides)
    return payload


def _setup_inquiries(tmp_path):
    db_path = tmp_path / "inquiry_case_manager.db"
    init_db(db_path)
    create_inquiry(
        _inquiry(status="未対応", subject="未対応案件", due_date=date(2026, 9, 10)),
        db_path=db_path,
    )
    create_inquiry(
        _inquiry(
            status="対応中",
            subject="対応中案件",
            customer_name="佐藤花子",
            company_name="",
            phone="",
            email="",
            notes="",
            due_date=date(2026, 9, 30),
        ),
        db_path=db_path,
    )
    create_inquiry(
        _inquiry(status="完了", subject="完了案件", assignee="佐藤", due_date=date(2026, 9, 10)),
        db_path=db_path,
    )
    return db_path


def _read_csv(csv_bytes):
    text = csv_bytes.decode("utf-8-sig")
    return list(csv.reader(io.StringIO(text)))


def _load_excel(excel_bytes):
    return load_workbook(io.BytesIO(excel_bytes))


def test_generate_inquiry_csv(tmp_path):
    db_path = _setup_inquiries(tmp_path)
    inquiries = search_inquiries(db_path=db_path, today=TODAY)

    csv_bytes = create_inquiry_csv_bytes(inquiries)

    assert csv_bytes.startswith(b"\xef\xbb\xbf")
    rows = _read_csv(csv_bytes)
    assert rows[0] == inquiry_export_headers()
    assert "問い合わせID" in rows[0]
    assert len(rows) - 1 == len(inquiries)


def test_generate_case_csv(tmp_path):
    db_path = _setup_inquiries(tmp_path)
    inquiries = search_inquiries(db_path=db_path, today=TODAY)
    for item in inquiries:
        create_case_from_inquiry(item["inquiry_id"], db_path=db_path, today=TODAY)
    cases = search_cases(db_path=db_path, today=TODAY)

    csv_bytes = create_case_csv_bytes(cases)

    assert csv_bytes.startswith(b"\xef\xbb\xbf")
    rows = _read_csv(csv_bytes)
    assert rows[0] == case_export_headers()
    assert "案件ID" in rows[0]
    assert len(rows) - 1 == len(cases)


def test_inquiry_overdue_and_empty_values(tmp_path):
    db_path = _setup_inquiries(tmp_path)
    inquiries = search_inquiries(db_path=db_path, today=TODAY)
    rows = prepare_inquiry_export_rows(inquiries)
    headers = inquiry_export_headers()
    overdue_index = headers.index("期限超過")
    notes_index = headers.index("備考")
    company_index = headers.index("会社名")

    overdue_rows = [row for row in rows if row[overdue_index] == OVERDUE_LABEL]
    assert len(overdue_rows) == 1
    completed = next(row for row in rows if row[headers.index("ステータス")] == "完了")
    assert completed[overdue_index] == ""

    empty_company = next(row for row in rows if row[headers.index("顧客名")] == "佐藤花子")
    assert empty_company[company_index] == ""
    assert empty_company[notes_index] == ""
    assert "None" not in empty_company
    assert "NaN" not in empty_company
    assert "NaT" not in [str(value) for value in empty_company]


def test_generate_inquiry_excel(tmp_path):
    db_path = _setup_inquiries(tmp_path)
    inquiries = search_inquiries(db_path=db_path, today=TODAY)

    excel_bytes = create_inquiry_excel_bytes(inquiries)
    workbook = _load_excel(excel_bytes)
    sheet = workbook.active

    assert sheet.title == INQUIRY_SHEET_NAME
    assert [cell.value for cell in sheet[1]] == inquiry_export_headers()
    assert all(cell.font.bold for cell in sheet[1])
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref
    assert sheet.max_row - 1 == len(inquiries)


def test_generate_case_excel(tmp_path):
    db_path = _setup_inquiries(tmp_path)
    inquiries = search_inquiries(db_path=db_path, today=TODAY)
    for item in inquiries:
        create_case_from_inquiry(item["inquiry_id"], db_path=db_path, today=TODAY)
    cases = search_cases(db_path=db_path, today=TODAY)

    excel_bytes = create_case_excel_bytes(cases)
    workbook = _load_excel(excel_bytes)
    sheet = workbook.active

    assert sheet.title == CASE_SHEET_NAME
    assert [cell.value for cell in sheet[1]] == case_export_headers()
    assert all(cell.font.bold for cell in sheet[1])
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref
    assert sheet.max_row - 1 == len(cases)


def test_excel_overdue_and_empty_cells(tmp_path):
    db_path = _setup_inquiries(tmp_path)
    inquiries = search_inquiries(db_path=db_path, today=TODAY)
    workbook = _load_excel(create_inquiry_excel_bytes(inquiries))
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    overdue_index = headers.index("期限超過") + 1
    company_index = headers.index("会社名") + 1
    notes_index = headers.index("備考") + 1
    customer_index = headers.index("顧客名") + 1

    overdue_values = [sheet.cell(row=index, column=overdue_index).value for index in range(2, sheet.max_row + 1)]
    assert overdue_values.count(OVERDUE_LABEL) == 1
    assert overdue_values.count(None) + overdue_values.count("") == 2

    for index in range(2, sheet.max_row + 1):
        if sheet.cell(row=index, column=customer_index).value == "佐藤花子":
            assert sheet.cell(row=index, column=company_index).value in (None, "")
            assert sheet.cell(row=index, column=notes_index).value in (None, "")


def test_inquiry_export_matches_filtered_search(tmp_path):
    db_path = _setup_inquiries(tmp_path)
    results = search_inquiries({"statuses": ["対応中"]}, db_path=db_path, today=TODAY)

    csv_rows = _read_csv(create_inquiry_csv_bytes(results))
    excel = _load_excel(create_inquiry_excel_bytes(results))

    assert len(results) == 1
    assert len(csv_rows) - 1 == len(results)
    assert excel.active.max_row - 1 == len(results)
    assert csv_rows[1][inquiry_export_headers().index("ステータス")] == "対応中"


def test_case_export_matches_filtered_search(tmp_path):
    db_path = _setup_inquiries(tmp_path)
    inquiries = search_inquiries(db_path=db_path, today=TODAY)
    for item in inquiries:
        create_case_from_inquiry(item["inquiry_id"], db_path=db_path, today=TODAY)
    results = search_cases({"priorities": ["高"]}, db_path=db_path, today=TODAY)

    csv_rows = _read_csv(create_case_csv_bytes(results))
    excel = _load_excel(create_case_excel_bytes(results))

    assert len(results) == 3
    assert len(csv_rows) - 1 == len(results)
    assert excel.active.max_row - 1 == len(results)


def test_generate_export_filename():
    now = datetime(2026, 9, 21, 8, 5, 6)

    assert generate_export_filename("inquiries", "csv", now=now) == "inquiries_20260921_080506.csv"
    assert generate_export_filename("cases", "xlsx", now=now) == "cases_20260921_080506.xlsx"


def test_empty_export_has_header_only():
    csv_rows = _read_csv(create_inquiry_csv_bytes([]))
    excel = _load_excel(create_inquiry_excel_bytes([]))

    assert csv_rows == [inquiry_export_headers()]
    assert excel.active.max_row == 1
    assert [cell.value for cell in excel.active[1]] == inquiry_export_headers()

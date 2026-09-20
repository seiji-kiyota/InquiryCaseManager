from datetime import date, datetime, time

from common.case_service import (
    create_case_from_inquiry,
    get_case_by_id,
    search_cases,
    update_case,
)
from common.dashboard_service import (
    count_cases_by_status,
    count_inquiries_by_status,
    get_case_kpis,
    get_inquiry_kpis,
)
from common.db import init_db
from common.export_service import (
    create_case_csv_bytes,
    create_case_excel_bytes,
    create_inquiry_csv_bytes,
    create_inquiry_excel_bytes,
)
from common.history_service import create_history, get_histories_by_target
from common.inquiry_service import (
    create_inquiry,
    get_inquiry_by_id,
    search_inquiries,
    update_inquiry,
)
from common.models import INQUIRY_TARGET_TYPE, STATUS_CHANGE_ACTION_TYPE
from openpyxl import load_workbook
import io


TODAY = date(2026, 9, 21)


def _inquiry(**overrides):
    payload = {
        "received_date": date(2026, 9, 10),
        "received_time": time(10, 0, 0),
        "customer_name": "山田太郎",
        "company_name": "株式会社サンプル",
        "phone": "03-1111-1111",
        "email": "taro@example.com",
        "channel": "電話",
        "category": "商品",
        "subject": "見積依頼",
        "description": "商品の見積をお願いします。",
        "priority": "高",
        "assignee": "山田",
        "due_date": date(2026, 9, 30),
        "status": "未対応",
        "notes": "",
    }
    payload.update(overrides)
    return payload


def _inquiry_update(existing, **overrides):
    payload = {
        "customer_name": existing["customer_name"],
        "company_name": existing["company_name"],
        "phone": existing["phone"],
        "email": existing["email"],
        "channel": existing["channel"],
        "category": existing["category"],
        "subject": existing["subject"],
        "description": existing["description"],
        "priority": existing["priority"],
        "assignee": existing["assignee"],
        "due_date": existing["due_date"],
        "status": existing["status"],
        "notes": existing["notes"],
    }
    payload.update(overrides)
    return payload


def _case_update(case, **overrides):
    payload = {
        "case_name": case["case_name"],
        "customer_name": case["customer_name"],
        "assignee": case["assignee"],
        "sub_assignee": case["sub_assignee"],
        "priority": case["priority"],
        "status": case["status"],
        "progress": case["progress"],
        "start_date": case["start_date"],
        "due_date": case["due_date"],
        "completed_date": case["completed_date"],
        "summary": case["summary"],
        "action_plan": case["action_plan"],
        "notes": case["notes"],
    }
    payload.update(overrides)
    return payload


def _csv_data_rows(csv_bytes):
    return len(csv_bytes.decode("utf-8-sig").splitlines()) - 1


def _excel_data_rows(excel_bytes):
    workbook = load_workbook(io.BytesIO(excel_bytes))
    return workbook.active.max_row - 1


def test_inquiry_to_case_completion_workflow(tmp_path):
    db_path = tmp_path / "inquiry_case_manager.db"
    init_db(db_path)

    inquiry_id = create_inquiry(_inquiry(), db_path=db_path)
    listed = search_inquiries(db_path=db_path, today=TODAY)
    found = search_inquiries({"keyword": "見積"}, db_path=db_path, today=TODAY)
    assert inquiry_id in [item["inquiry_id"] for item in listed]
    assert len(found) == 1

    existing = get_inquiry_by_id(inquiry_id, db_path=db_path)
    updated_id = update_inquiry(
        inquiry_id,
        _inquiry_update(existing, description="内容を更新しました。", status="対応中"),
        db_path=db_path,
        now=datetime(2026, 9, 12, 11, 0, 0),
    )
    assert updated_id == inquiry_id
    assert get_inquiry_by_id(inquiry_id, db_path=db_path)["status"] == "対応中"

    history_id = create_history(
        {
            "target_type": INQUIRY_TARGET_TYPE,
            "target_id": inquiry_id,
            "action_datetime": datetime(2026, 9, 12, 11, 30, 0),
            "operator": "山田",
            "action_type": "電話",
            "action_detail": "内容を確認しました。",
        },
        db_path=db_path,
    )
    histories = get_histories_by_target(INQUIRY_TARGET_TYPE, inquiry_id, db_path=db_path)
    assert history_id is not None
    assert any(item["action_type"] == "電話" for item in histories)
    assert any(
        item["action_type"] == STATUS_CHANGE_ACTION_TYPE and item["new_status"] == "対応中"
        for item in histories
    )

    case_id = create_case_from_inquiry(inquiry_id, db_path=db_path, today=TODAY)
    cases = search_cases(db_path=db_path, today=TODAY)
    assert case_id == "CASE-202609-0001"
    assert len(cases) == 1

    case = get_case_by_id(case_id, db_path=db_path, today=TODAY)
    update_case(
        case_id,
        _case_update(case, status="対応中", progress=40),
        db_path=db_path,
        today=TODAY,
    )
    case = get_case_by_id(case_id, db_path=db_path, today=TODAY)
    update_case(
        case_id,
        _case_update(case, status="完了"),
        db_path=db_path,
        today=TODAY,
    )
    completed = get_case_by_id(case_id, db_path=db_path, today=TODAY)
    assert completed["status"] == "完了"
    assert completed["progress"] == 100
    assert completed["completed_date"] == TODAY.isoformat()

    inquiry_kpis = get_inquiry_kpis(db_path=db_path, today=TODAY)
    case_kpis = get_case_kpis(db_path=db_path, today=TODAY)
    inquiry_status = count_inquiries_by_status(db_path=db_path)
    case_status = count_cases_by_status(db_path=db_path)
    assert inquiry_kpis["total"] == sum(inquiry_status.values()) == 1
    assert inquiry_kpis["対応中"] == 1
    assert inquiry_kpis["completion_rate"] == 0.0
    assert case_kpis["total"] == sum(case_status.values()) == 1
    assert case_kpis["完了"] == 1
    assert case_kpis["completion_rate"] == 100.0

    inquiries = search_inquiries(db_path=db_path, today=TODAY)
    cases = search_cases(db_path=db_path, today=TODAY)
    assert _csv_data_rows(create_inquiry_csv_bytes(inquiries)) == len(inquiries)
    assert _excel_data_rows(create_inquiry_excel_bytes(inquiries)) == len(inquiries)
    assert _csv_data_rows(create_case_csv_bytes(cases)) == len(cases)
    assert _excel_data_rows(create_case_excel_bytes(cases)) == len(cases)


def test_filtered_export_matches_search_counts(tmp_path):
    db_path = tmp_path / "inquiry_case_manager.db"
    init_db(db_path)
    create_inquiry(_inquiry(status="未対応", subject="未対応A"), db_path=db_path)
    create_inquiry(
        _inquiry(status="対応中", subject="対応中B", customer_name="佐藤花子"),
        db_path=db_path,
    )
    create_inquiry(_inquiry(status="完了", subject="完了C"), db_path=db_path)
    first = search_inquiries({"statuses": ["対応中"]}, db_path=db_path, today=TODAY)
    create_case_from_inquiry(first[0]["inquiry_id"], db_path=db_path, today=TODAY)
    cases = search_cases({"statuses": ["未着手"]}, db_path=db_path, today=TODAY)

    assert len(first) == 1
    assert _csv_data_rows(create_inquiry_csv_bytes(first)) == 1
    assert _excel_data_rows(create_inquiry_excel_bytes(first)) == 1
    assert len(cases) == 1
    assert _csv_data_rows(create_case_csv_bytes(cases)) == 1
    assert _excel_data_rows(create_case_excel_bytes(cases)) == 1

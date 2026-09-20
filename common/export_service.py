"""問い合わせ・案件のCSV / Excel出力。"""

import csv
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

OVERDUE_LABEL = "期限超過"
INQUIRY_SHEET_NAME = "問い合わせ一覧"
CASE_SHEET_NAME = "案件一覧"

INQUIRY_EXPORT_COLUMNS = (
    ("問い合わせID", "inquiry_id"),
    ("受付日", "received_date"),
    ("受付時刻", "received_time"),
    ("顧客名", "customer_name"),
    ("会社名", "company_name"),
    ("電話番号", "phone"),
    ("メールアドレス", "email"),
    ("受付方法", "channel"),
    ("カテゴリ", "category"),
    ("件名", "subject"),
    ("問い合わせ内容", "description"),
    ("優先度", "priority"),
    ("担当者", "assignee"),
    ("対応期限", "due_date"),
    ("ステータス", "status"),
    ("備考", "notes"),
    ("期限超過", "overdue"),
    ("作成日時", "created_at"),
    ("更新日時", "updated_at"),
)

CASE_EXPORT_COLUMNS = (
    ("案件ID", "case_id"),
    ("元問い合わせID", "inquiry_id"),
    ("案件名", "case_name"),
    ("顧客名", "customer_name"),
    ("担当者", "assignee"),
    ("副担当者", "sub_assignee"),
    ("優先度", "priority"),
    ("ステータス", "status"),
    ("進捗率", "progress"),
    ("開始日", "start_date"),
    ("期限", "due_date"),
    ("完了日", "completed_date"),
    ("概要", "summary"),
    ("対応方針", "action_plan"),
    ("備考", "notes"),
    ("期限超過", "overdue"),
    ("作成日時", "created_at"),
    ("更新日時", "updated_at"),
)

_EMPTY_MARKERS = {"", "none", "nan", "nat"}


def _export_value(item, key):
    value = item.get(key)
    if key == "overdue":
        if value is True or value == OVERDUE_LABEL:
            return OVERDUE_LABEL
        return ""
    if value is None:
        return ""
    if isinstance(value, str):
        if value.strip().lower() in _EMPTY_MARKERS:
            return ""
        return value
    return value


def inquiry_export_headers():
    return [header for header, _key in INQUIRY_EXPORT_COLUMNS]


def case_export_headers():
    return [header for header, _key in CASE_EXPORT_COLUMNS]


def prepare_inquiry_export_rows(inquiries):
    """問い合わせを出力用の行データへ整形する。"""
    return [
        [_export_value(item, key) for _header, key in INQUIRY_EXPORT_COLUMNS]
        for item in inquiries or []
    ]


def prepare_case_export_rows(cases):
    """案件を出力用の行データへ整形する。"""
    return [
        [_export_value(item, key) for _header, key in CASE_EXPORT_COLUMNS]
        for item in cases or []
    ]


def generate_export_filename(prefix, extension, now=None):
    """出力日時付きのファイル名を返す。"""
    timestamp = (now or datetime.now()).strftime("%Y%m%d_%H%M%S")
    return f"{prefix}_{timestamp}.{extension.lstrip('.')}"


def create_csv_bytes(headers, rows):
    """UTF-8 BOM付きCSVのバイト列を返す。"""
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(headers)
    for row in rows:
        writer.writerow(["" if value is None else value for value in row])
    return buffer.getvalue().encode("utf-8-sig")


def create_excel_bytes(headers, rows, sheet_name):
    """ヘッダー太字・1行目固定・フィルター付きExcelのバイト列を返す。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(["" if value is None else value for value in row])

    last_column = get_column_letter(len(headers))
    last_row = max(sheet.max_row, 1)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{last_column}{last_row}"

    for index, column_cells in enumerate(sheet.columns, start=1):
        max_length = 0
        for cell in column_cells:
            text = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(text))
        sheet.column_dimensions[get_column_letter(index)].width = min(max(max_length + 2, 8), 40)

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def create_inquiry_csv_bytes(inquiries):
    return create_csv_bytes(inquiry_export_headers(), prepare_inquiry_export_rows(inquiries))


def create_case_csv_bytes(cases):
    return create_csv_bytes(case_export_headers(), prepare_case_export_rows(cases))


def create_inquiry_excel_bytes(inquiries):
    return create_excel_bytes(
        inquiry_export_headers(),
        prepare_inquiry_export_rows(inquiries),
        INQUIRY_SHEET_NAME,
    )


def create_case_excel_bytes(cases):
    return create_excel_bytes(
        case_export_headers(),
        prepare_case_export_rows(cases),
        CASE_SHEET_NAME,
    )
